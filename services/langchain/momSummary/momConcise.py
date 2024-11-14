from utilservice import *
from fastapi import HTTPException, status
from langchain_community.document_loaders import Docx2txtLoader
from langchain_community.llms.ollama import Ollama
from langchain.schema import Document
import asyncio
import os
import docx
from time import time
from services.langchain.langchainUtils.singleton import OllamaSingleton


ollama_model = OllamaSingleton.get_instance()


concise_summary_prompt = """

You are an advanced summarization AI. Your task is to create a concise summary from the provided transcript. The summary should be one cohesive paragraph.
The summary should encapsulate all key points while maintaining brevity and clarity. Each point should be distilled into a concise form, ensuring that no essential information is omitted.

Instructions:
1. Read the provided transcript excerpt carefully.
2. Identify and extract all significant points, arguments, and conclusions from the text.
3. Summarize each identified point in one or two sentences, ensuring it reflects the essence of the original content.
4. Organize the summarized points in a logical order, maintaining coherence.
5. Ensure the final output is a single paragraph that includes all the concise points without unnecessary elaboration.

=========================================================================================================================================

TRANSCRIPT :

"""

executive_summary_prompt = """

Generate a concise, professional executive summary based on the provided transcript. The summary should be **one cohesive paragraph**. Capture all key points discussed, decisions made, and any next steps or actions agreed upon. Be clear, concise, and written in professional business English. Use formal language and avoid colloquialisms or casual expressions. Be specific and avoid vague statements. Do not include unnecessary details or placeholders. Reflect the actual points discussed in the meeting. Please generate text without any markdown headers or hashes.

"""

meeting_notes_prompt = """
Provide detailed meeting notes based on the transcript. The notes should clearly outline all the topics discussed, specific feedback or comments given, and any comparisons made. Also include additional suggestions, discussions, future plans, and priorities. Structure the notes by topic, ensuring each section is clearly organized using dash (-) points for clarity. Do not include the topic name headings like Action items.

Template:

*Topic Name*
- [Summarize key point1 discussed under this topic]
- [Summarize key point2 discussed under this topic]

*Another Topic Name*
- [Summarize key point1 discussed under this topic]
- [Summarize key point2 discussed under this topic]

... (continue for all topics discussed in the meeting)

(Note: Ensure each topic is described concisely and avoid repetitive phrases. Maintain a continuous flow of information as seen from the transcript. Please generate text without any markdown headers or hashes.)

"""

action_items_prompt = """

Based on the provided transcript, generate a list of only the most important and high-priority action items discussed during the meeting. Focus on key decisions and tasks that have a significant impact on the project or require urgent attention. Avoid including minor tasks, repetitive actions, or general suggestions. Summarize the actions concisely and clearly, while keeping the language professional and specific.

Format the action items as follows:
- Action Item 1: [describe the key task/decision],

General guidelines for the outputs:
- Ensure all content is based solely on the provided transcript.
- Do not fabricate content or include placeholders.
"""

short_summary_prompt = """

Based on the following full meeting summary, generate a concise short summary. Limit the summary to 10-12 bullet points.
Ensure each point is unique, concise, and avoids repetition.

Full Meeting Summary:

"""

async def load_and_extract_document_content(data_path, filename):
    print(f"Loading document content from: {data_path}/{filename}")  # Path of the document being loaded
    doc = docx.Document(f"{data_path}/{filename}")
    document_objects = []

    for paragraph in doc.paragraphs:
        paragraph_text = paragraph.text.strip()
        if paragraph_text:
            document_objects.append(paragraph_text)

    print(f"Extracted {len(document_objects)} paragraphs from the document.")  # Number of paragraphs extracted
    return document_objects

async def create_concise_summaries(paragraphs):
    print(f"Creating concise summaries for {len(paragraphs)} paragraphs.")  # Number of paragraphs to summarize
    combined_summary = []  # Use a list to collect summaries

    for paragraph in paragraphs:
        prompt = concise_summary_prompt + paragraph
        print(f"Input to concise summary model: {prompt}")  # Print full input prompt
        summary = await asyncio.to_thread(ollama_model.invoke, prompt)
        # Remove the unwanted phrase if it exists and strip whitespace
        cleaned_summary = summary.replace("Here is a concise summary of the transcript in one paragraph:", "").strip()
        combined_summary.append(cleaned_summary)  # Append the cleaned summary to the list
        print(f"Generated concise summary: {cleaned_summary}")  # Print the generated summary for the paragraph

    # Combine all summaries into a single paragraph
    return " ".join(combined_summary)  # Join the summaries into one paragraph

async def create_meeting_summary(concise_summary):
    print("Creating meeting summary based on the concise summary.")  # Status of meeting summary creation
    mom_sections = ["Executive_Summary", "Meeting_Notes", "Action_Item"]
    mom_sections_prompts = [executive_summary_prompt, meeting_notes_prompt, action_items_prompt]
    mom_summary = {}

    for i, prompt in enumerate(mom_sections_prompts):
        prompt = prompt + concise_summary
        print(f"Input for {mom_sections[i]} model: {prompt}")  # Print the full input for each section
        summary_response = await asyncio.to_thread(ollama_model.invoke, prompt)
        mom_summary[mom_sections[i]] = summary_response.strip()
        print(f"{mom_sections[i]} generated: {summary_response.strip()}")  # Print the generated summary for each section

    return mom_summary

async def create_short_summaries(full_mom_summary):
    print("Creating short summary based on the full meeting summary.")  # Status of short summary creation
    prompt = short_summary_prompt + str(full_mom_summary)
    print(f"Input for short summary model: {prompt}")  # Print the full input for the short summary
    summary = await asyncio.to_thread(ollama_model.invoke, prompt)
    print(f"Short summary generated: {summary}")  # Print the generated short summary

    return summary

async def generate_mom_summary(llm_model, data_path, filename, recording_id):
    start_time = time()
    print(f"Generating MOM summary for recordingID: {recording_id}...")  # Status of summary generation

    paragraphs = await load_and_extract_document_content(data_path, filename)
    concise_text_summary = await create_concise_summaries(paragraphs)
    full_mom_summary = await create_meeting_summary(concise_text_summary)
    short_mom_summary = await create_short_summaries(full_mom_summary)

    execution_time = f"{time() - start_time:.2f}"
    print(f"MOM Summary generated successfully for recordingID: {recording_id} in {execution_time} seconds.", "\n")

    final_summary_object = {
        "MOM_SUMMARY": full_mom_summary,
        "SHORT_SUMMARY": short_mom_summary
    }

    return final_summary_object
